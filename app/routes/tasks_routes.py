"""Tasks UI and JSON API, migrated from the supplied SQLite application."""

from datetime import datetime
import sqlite3
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from flask import Blueprint, abort, current_app, jsonify, make_response, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.utils import secure_filename

from app.models import tasks_model as store


tasks_bp = Blueprint('tasks', __name__)
MAX_FILE_SIZE = 50 * 1024 * 1024


def _owner(): return session.get('user_id')
def _upload_dir():
    """Keep uploads beside the SQLite database, not relative to the service CWD."""
    directory = Path(current_app.instance_path) / 'uploads' / 'tasks'
    directory.mkdir(parents=True, exist_ok=True)
    return directory
def _auth():
    if not _owner(): return None
    return _owner()
def _task_data(data):
    title = str(data.get('title', '')).strip()
    if not title: raise ValueError('Judul tugas wajib diisi.')
    description = str(data.get('description', '')).strip()
    if len(description) > 10000: raise ValueError('Notes may contain at most 10,000 characters.')
    priority = data.get('priority', 'medium')
    return {'title': title, 'description': description, 'category': str(data.get('category', 'Personal')).strip() or 'Personal', 'due_date': data.get('due_date') or None, 'priority': priority if priority in {'low','medium','high'} else 'medium'}
def _save_upload(file):
    name = secure_filename(file.filename or '')
    if not name: raise ValueError('Nama file tidak valid.')
    file.stream.seek(0, 2); size=file.stream.tell(); file.stream.seek(0)
    if size > MAX_FILE_SIZE: raise OverflowError
    stored_name = f"{datetime.now():%Y%m%d%H%M%S%f}_{uuid4().hex}_{name}"
    file.save(_upload_dir() / stored_name); return name, stored_name, size
def _remove(names):
    for name in names:
        (_upload_dir() / Path(name).name).unlink(missing_ok=True)
def _json(record): return store.serialize(record)


@tasks_bp.get('/user/tasks')
def index():
    if not _auth(): return redirect(url_for('auth.login'))
    # The form contains a session-bound CSRF token.  Never let a browser reuse
    # a page cached before an application restart or a new user session.
    response = make_response(render_template('user/tasks.html', app_version='v1.3.0', last_updated='2026-09-05'))
    response.headers['Cache-Control'] = 'no-store, max-age=0'
    response.headers['Pragma'] = 'no-cache'
    return response


@tasks_bp.get('/api/tasks')
def get_tasks():
    owner=_auth()
    if not owner: return jsonify(error='Login diperlukan.'), 401
    try: page=max(1,int(request.args.get('page',1)))
    except ValueError: page=1
    records,total=store.list_tasks(owner, request.args.get('status','all'), request.args.get('search','').strip(), request.args.get('date_from','').strip(), request.args.get('date_to','').strip(), page)
    total_pages=max(1,(total+4)//5); page=min(page,total_pages)
    if page != request.args.get('page', page) and total: records,_=store.list_tasks(owner, request.args.get('status','all'), request.args.get('search','').strip(), request.args.get('date_from','').strip(), request.args.get('date_to','').strip(), page)
    payload={'items':[_json(item) for item in records],'page':page,'per_page':5,'total':total,'total_pages':total_pages}
    return jsonify(payload if 'page' in request.args else payload['items'])


@tasks_bp.get('/api/stats')
def stats():
    owner=_auth()
    if not owner: return jsonify(error='Login diperlukan.'),401
    return jsonify(store.stats(owner,request.args.get('search','').strip(),request.args.get('date_from','').strip(),request.args.get('date_to','').strip()))


@tasks_bp.post('/api/tasks')
def create_task():
    owner=_auth()
    if not owner: return jsonify(error='Login diperlukan.'),401
    try: task=store.insert_task(owner,_task_data(request.get_json() or {}))
    except ValueError as error: return jsonify(error=str(error)),400
    except sqlite3.Error as error:
        current_app.logger.exception("Unable to save Task: %s", error)
        if 'locked' in str(error).lower() or 'busy' in str(error).lower():
            return jsonify(error='Task could not be saved because the SQLite database is busy. Please try again in a moment.'), 503
        return jsonify(error='Task could not be saved because the local Tasks database schema needs to be updated.'), 500
    return jsonify(_json(task)),201


@tasks_bp.get('/api/tasks/<int:task_id>')
def get_task(task_id):
    owner=_auth(); task=store.get_task(owner,task_id) if owner else None
    return jsonify(_json(task)) if task else (jsonify(error='Tugas tidak ditemukan.'),404)


@tasks_bp.patch('/api/tasks/<int:task_id>')
def update_task(task_id):
    owner=_auth()
    if not owner:return jsonify(error='Login diperlukan.'),401
    data=request.get_json() or {}
    if 'title' in data and not str(data['title']).strip(): return jsonify(error='Judul tugas wajib diisi.'),400
    if 'description' in data and len(str(data['description'])) > 10000: return jsonify(error='Notes may contain at most 10,000 characters.'),400
    task=store.update_task(owner,task_id,data)
    return jsonify(_json(task)) if task else (jsonify(error='Tugas tidak ditemukan atau tidak ada perubahan.'),404)


@tasks_bp.delete('/api/tasks/<int:task_id>')
def delete_task(task_id):
    owner=_auth(); files=store.delete_task(owner,task_id) if owner else None
    if files is None:return jsonify(error='Tugas tidak ditemukan.'),404
    _remove(files); return '',204


@tasks_bp.post('/api/tasks/<int:task_id>/attachments')
def upload_attachment(task_id):
    owner=_auth(); file=request.files.get('file')
    if not owner:return jsonify(error='Login diperlukan.'),401
    if not file or not file.filename:return jsonify(error='Pilih file untuk diunggah.'),400
    try: original,stored,size=_save_upload(file)
    except ValueError as error:return jsonify(error=str(error)),400
    except OverflowError:return jsonify(error='Ukuran file maksimal 50 MB.'),413
    except OSError:return jsonify(error='Lampiran tidak dapat disimpan. Periksa izin folder instance/uploads/tasks.'),500
    attachment_id=store.add_attachment(owner,task_id,original,stored,size)
    if not attachment_id:_remove([stored]);return jsonify(error='Tugas tidak ditemukan.'),404
    return jsonify(id=attachment_id,original_name=original,size=size),201


@tasks_bp.get('/api/attachments/<int:attachment_id>/download')
def download_attachment(attachment_id):
    item=store.attachment(_owner(),attachment_id) if _auth() else None
    if not item:abort(404)
    return send_file(_upload_dir()/item['stored_name'],as_attachment=True,download_name=item['original_name'])


@tasks_bp.delete('/api/attachments/<int:attachment_id>')
def delete_attachment(attachment_id):
    name=store.delete_attachment(_owner(),attachment_id) if _auth() else None
    if not name:return jsonify(error='Lampiran tidak ditemukan.'),404
    _remove([name]);return '',204


@tasks_bp.get('/api/tasks/<int:task_id>/subtasks/<int:subtask_id>')
def get_subtask(task_id,subtask_id):
    subtask=store.get_subtask(_owner(),task_id,subtask_id) if _auth() else None
    return jsonify(_json(subtask)) if subtask else (jsonify(error='Subtugas tidak ditemukan.'),404)


@tasks_bp.post('/api/tasks/<int:task_id>/subtasks')
@tasks_bp.patch('/api/tasks/<int:task_id>/subtasks/<int:subtask_id>')
def save_subtask(task_id,subtask_id=None):
    owner=_auth(); data=request.get_json() or {}; name=str(data.get('name','')).strip()
    if not owner:return jsonify(error='Login diperlukan.'),401
    if not name:return jsonify(error='Nama subtugas wajib diisi.'),400
    try:
        item=store.save_subtask(owner,task_id,{'name':name,'due_date':data.get('due_date') or None,'notes':str(data.get('notes','')).strip()},subtask_id)
    except sqlite3.Error as error:
        current_app.logger.exception("Unable to save Step: %s", error)
        if 'locked' in str(error).lower() or 'busy' in str(error).lower():
            return jsonify(error='Step storage is busy. Please try saving again.'), 503
        return jsonify(error='Step could not be saved because the local Tasks database schema needs to be updated.'), 500
    if not item:return jsonify(error='Subtugas tidak ditemukan.'),404
    return jsonify(_json(item)), 200 if subtask_id else 201


@tasks_bp.delete('/api/tasks/<int:task_id>/subtasks/<int:subtask_id>')
def delete_subtask(task_id,subtask_id):
    files=store.delete_subtask(_owner(),task_id,subtask_id) if _auth() else None
    if files is None:return jsonify(error='Subtugas tidak ditemukan.'),404
    _remove(files);return '',204


@tasks_bp.post('/api/tasks/<int:task_id>/subtasks/<int:subtask_id>/attachments')
def upload_subtask_attachment(task_id,subtask_id):
    owner=_auth(); file=request.files.get('file')
    if not owner:return jsonify(error='Login diperlukan.'),401
    if not file or not file.filename:return jsonify(error='Pilih file untuk diunggah.'),400
    try: original,stored,size=_save_upload(file)
    except ValueError as error:return jsonify(error=str(error)),400
    except OverflowError:return jsonify(error='Ukuran file maksimal 50 MB.'),413
    except OSError:return jsonify(error='Lampiran tidak dapat disimpan. Periksa izin folder instance/uploads/tasks.'),500
    attachment_id=store.add_attachment(owner,task_id,original,stored,size,subtask_id)
    if not attachment_id:_remove([stored]);return jsonify(error='Subtugas tidak ditemukan.'),404
    return jsonify(id=attachment_id,original_name=original,size=size),201


@tasks_bp.get('/api/subtask-attachments/<int:attachment_id>/download')
def download_subtask_attachment(attachment_id):
    item=store.attachment(_owner(),attachment_id,True) if _auth() else None
    if not item:abort(404)
    return send_file(_upload_dir()/item['stored_name'],as_attachment=True,download_name=item['original_name'])


@tasks_bp.delete('/api/subtask-attachments/<int:attachment_id>')
def delete_subtask_attachment(attachment_id):
    name=store.delete_attachment(_owner(),attachment_id,True) if _auth() else None
    if not name:return jsonify(error='Lampiran tidak ditemukan.'),404
    _remove([name]);return '',204


@tasks_bp.get('/api/tasks/export')
def export_tasks():
    owner=_auth()
    if not owner:return jsonify(error='Login diperlukan.'),401
    records,total=store.list_tasks(owner,request.args.get('status','all'),request.args.get('search','').strip(),request.args.get('date_from','').strip(),request.args.get('date_to','').strip(),1,701)
    if total>700:return jsonify(error='Ekspor dibatasi maksimal 700 task. Persempit filter lalu coba lagi.'),422
    workbook=Workbook(); sheet=workbook.active; sheet.title='Tasks'; sheet.append(['ID','Nama Task','Status','Kategori','Prioritas','Due Date','Catatan','Dibuat Pada'])
    for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1F4E78')
    for task in records: sheet.append([task['id'],task['title'],'Canceled' if task['canceled'] else 'Deployed' if task['completed'] else 'Active',task['category'],task['priority'],task['due_date'] or '',task['description'] or '',task['created_at']])
    output=BytesIO();workbook.save(output);output.seek(0)
    return send_file(output,as_attachment=True,download_name=f"tasks_export_{datetime.now():%Y%m%d_%H%M%S}.xlsx",mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
